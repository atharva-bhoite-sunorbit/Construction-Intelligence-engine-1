"""
Export Engine for Geotechnical Intelligence Reports.
Generates comprehensive engineering PDFs and multi-tab Excel workbooks.
"""
from __future__ import annotations

import io
from typing import Any, Dict, List
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from reportlab.lib.pagesizes import letter
from reportlab.lib import colors
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, KeepTogether
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle

from backend.app.services.geotech_intel.fields import NOT_IN_REPORT


def export_geotechnical_excel(report_data: Dict[str, Any]) -> bytes:
    """Generates a professional multi-sheet Excel workbook from geotechnical report data."""
    wb = openpyxl.Workbook()
    wb.remove(wb.active)  # Remove default sheet

    header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    sec_font = Font(name="Segoe UI", size=13, bold=True, color="1E3A8A")
    cell_font = Font(name="Segoe UI", size=10)
    thin_border = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )

    # 1. Summary Sheet
    ws_sum = wb.create_sheet(title="Executive Summary")
    ws_sum.column_dimensions['A'].width = 30
    ws_sum.column_dimensions['B'].width = 45
    ws_sum.column_dimensions['C'].width = 18
    ws_sum.column_dimensions['D'].width = 14

    ws_sum.append(["GEOTECHNICAL INVESTIGATION INTELLIGENCE SUMMARY", "", "", ""])
    ws_sum.cell(1, 1).font = Font(name="Segoe UI", size=14, bold=True, color="1E3A8A")
    ws_sum.append([])

    headers = ["Parameter", "Extracted Value", "Source Type", "Confidence"]
    ws_sum.append(headers)
    for col_idx in range(1, 5):
        c = ws_sum.cell(3, col_idx)
        c.font = header_font
        c.fill = header_fill
        c.alignment = Alignment(horizontal="center", vertical="center")

    proj = report_data.get("project_information", {})
    summary_rows = [
        ("Project Name", proj.get("project_name", {}).get("display", NOT_IN_REPORT), proj.get("project_name", {}).get("source_type", "REPORT"), proj.get("project_name", {}).get("confidence", "HIGH")),
        ("Client", proj.get("client", {}).get("display", NOT_IN_REPORT), proj.get("client", {}).get("source_type", "REPORT"), proj.get("client", {}).get("confidence", "HIGH")),
        ("Location", proj.get("location", {}).get("display", NOT_IN_REPORT), proj.get("location", {}).get("source_type", "REPORT"), proj.get("location", {}).get("confidence", "HIGH")),
        ("Building Floors", proj.get("number_of_floors", {}).get("display", NOT_IN_REPORT), proj.get("number_of_floors", {}).get("source_type", "REPORT"), proj.get("number_of_floors", {}).get("confidence", "HIGH")),
        ("Consultant", proj.get("consultant", {}).get("display", NOT_IN_REPORT), proj.get("consultant", {}).get("source_type", "REPORT"), proj.get("consultant", {}).get("confidence", "HIGH")),
        ("Boreholes Count", report_data.get("investigation_information", {}).get("number_of_boreholes", {}).get("display", NOT_IN_REPORT), "REPORT", "HIGH"),
        ("Max Net Bearing Capacity", report_data.get("foundation_recommendations", {}).get("net_allowable_bearing_capacity", {}).get("display", NOT_IN_REPORT), "REPORT", "HIGH"),
        ("Max Settlement", report_data.get("foundation_recommendations", {}).get("maximum_settlement", {}).get("display", NOT_IN_REPORT), "REPORT", "HIGH"),
        ("Groundwater Depth", report_data.get("groundwater_analysis", {}).get("observed_depth", {}).get("display", NOT_IN_REPORT), "REPORT", "HIGH"),
    ]

    for row_idx, r in enumerate(summary_rows, start=4):
        ws_sum.append(list(r))
        for col_idx in range(1, 5):
            c = ws_sum.cell(row_idx, col_idx)
            c.font = cell_font
            c.border = thin_border

    # 2. Borehole Database Sheet
    ws_bh = wb.create_sheet(title="Borehole Database")
    ws_bh.append(["Borehole ID", "CWR Depth (m)", "Hard Rock Depth (m)", "Termination Depth (m)", "Groundwater Depth", "Source Sheet"])
    for col_idx in range(1, 7):
        c = ws_bh.cell(1, col_idx)
        c.font = header_font
        c.fill = header_fill

    bhs = report_data.get("boreholes", [])
    for b in bhs:
        ws_bh.append([
            b.get("borehole_id"),
            b.get("cwr_depth"),
            b.get("hard_rock_depth"),
            b.get("termination_depth"),
            b.get("groundwater_depth"),
            b.get("source_reference")
        ])

    # 3. Risks Sheet
    ws_risk = wb.create_sheet(title="Risk Engine")
    ws_risk.append(["Risk Title", "Severity", "Category", "Reason", "Recommended Action"])
    for col_idx in range(1, 6):
        c = ws_risk.cell(1, col_idx)
        c.font = header_font
        c.fill = header_fill

    risks = report_data.get("risks", [])
    for r in risks:
        ws_risk.append([
            r.get("risk_title"),
            r.get("severity"),
            r.get("category"),
            r.get("reason"),
            r.get("recommended_action")
        ])

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def export_geotechnical_pdf(report_data: Dict[str, Any]) -> bytes:
    """Generates an executive engineering PDF report using ReportLab."""
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=letter, rightMargin=36, leftMargin=36, topMargin=36, bottomMargin=36)
    styles = getSampleStyleSheet()

    title_style = ParagraphStyle(
        name="RepTitle",
        fontName="Helvetica-Bold",
        fontSize=18,
        leading=22,
        textColor=colors.HexColor("#1E3A8A"),
        spaceAfter=12
    )
    h2_style = ParagraphStyle(
        name="RepH2",
        fontName="Helvetica-Bold",
        fontSize=13,
        leading=16,
        textColor=colors.HexColor("#0F172A"),
        spaceBefore=14,
        spaceAfter=6
    )
    body_style = ParagraphStyle(
        name="RepBody",
        fontName="Helvetica",
        fontSize=9,
        leading=13,
        textColor=colors.HexColor("#334155")
    )
    body_bold = ParagraphStyle(
        name="RepBodyBold",
        fontName="Helvetica-Bold",
        fontSize=9,
        leading=13,
        textColor=colors.HexColor("#0F172A")
    )

    story = []

    # Title Banner
    story.append(Paragraph("GEOTECHNICAL INVESTIGATION INTELLIGENCE REPORT", title_style))
    story.append(Paragraph("Construction Intelligence Engine — Verified Source-Grounded Extraction", body_style))
    story.append(Spacer(1, 10))

    # Project Information Table
    proj = report_data.get("project_information", {})
    proj_table_data = [
        [Paragraph("Project Name", body_bold), Paragraph(proj.get("project_name", {}).get("display", NOT_IN_REPORT), body_style),
         Paragraph("Client", body_bold), Paragraph(proj.get("client", {}).get("display", NOT_IN_REPORT), body_style)],
        [Paragraph("Location", body_bold), Paragraph(proj.get("location", {}).get("display", NOT_IN_REPORT), body_style),
         Paragraph("Structure", body_bold), Paragraph(proj.get("building_configuration", {}).get("display", NOT_IN_REPORT), body_style)],
        [Paragraph("Consultant", body_bold), Paragraph(proj.get("consultant", {}).get("display", NOT_IN_REPORT), body_style),
         Paragraph("Investigation Date", body_bold), Paragraph(proj.get("investigation_date", {}).get("display", NOT_IN_REPORT), body_style)]
    ]
    t_proj = Table(proj_table_data, colWidths=[100, 170, 100, 170])
    t_proj.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, -1), colors.HexColor("#F8FAFC")),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor("#CBD5E1")),
        ('TOPPADDING', (0, 0), (-1, -1), 5),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 5),
    ]))
    story.append(t_proj)
    story.append(Spacer(1, 12))

    # Key Geotechnical Parameters Table
    story.append(Paragraph("Key Governing Geotechnical Parameters", h2_style))
    found = report_data.get("foundation_recommendations", {})
    rock = report_data.get("rock_analysis", {})
    gw = report_data.get("groundwater_analysis", {})

    params_data = [
        [Paragraph("Parameter", body_bold), Paragraph("Extracted Value", body_bold), Paragraph("Source Type", body_bold), Paragraph("Confidence", body_bold)],
        [Paragraph("Allowable Bearing Capacity", body_style), Paragraph(found.get("net_allowable_bearing_capacity", {}).get("display", NOT_IN_REPORT), body_style), Paragraph("REPORT", body_style), Paragraph("HIGH", body_style)],
        [Paragraph("Maximum Settlement", body_style), Paragraph(found.get("maximum_settlement", {}).get("display", NOT_IN_REPORT), body_style), Paragraph("REPORT", body_style), Paragraph("HIGH", body_style)],
        [Paragraph("Modulus of Subgrade Reaction", body_style), Paragraph(found.get("subgrade_reaction_modulus", {}).get("display", NOT_IN_REPORT), body_style), Paragraph("REPORT", body_style), Paragraph("HIGH", body_style)],
        [Paragraph("Observed Groundwater Depth", body_style), Paragraph(gw.get("observed_depth", {}).get("display", NOT_IN_REPORT), body_style), Paragraph("REPORT", body_style), Paragraph("HIGH", body_style)],
        [Paragraph("Rock Compressive Strength", body_style), Paragraph(rock.get("compressive_strength", {}).get("display", NOT_IN_REPORT), body_style), Paragraph("REPORT", body_style), Paragraph("HIGH", body_style)],
        [Paragraph("Equivalent UCS (MPa)", body_style), Paragraph(rock.get("compressive_strength_mpa_equivalent", {}).get("display", NOT_IN_REPORT), body_style), Paragraph("CALCULATED", body_style), Paragraph("HIGH", body_style)],
        [Paragraph("Core Recovery Range", body_style), Paragraph(rock.get("core_recovery", {}).get("display", NOT_IN_REPORT), body_style), Paragraph("REPORT", body_style), Paragraph("HIGH", body_style)],
        [Paragraph("RQD Range", body_style), Paragraph(rock.get("rqd", {}).get("display", NOT_IN_REPORT), body_style), Paragraph("REPORT", body_style), Paragraph("HIGH", body_style)],
    ]
    t_params = Table(params_data, colWidths=[160, 180, 100, 100])
    t_params.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor("#1E3A8A")),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor("#E2E8F0")),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor("#F8FAFC")]),
        ('TOPPADDING', (0, 0), (-1, -1), 4),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
    ]))
    story.append(t_params)
    story.append(Spacer(1, 12))

    # Independent Borehole Schedule Table
    story.append(Paragraph("Independent Borehole Schedule (Table A)", h2_style))
    bhs = report_data.get("boreholes", [])
    bh_headers = [Paragraph("Borehole", body_bold), Paragraph("CWR Depth", body_bold), Paragraph("Hard Rock Depth", body_bold), Paragraph("Termination", body_bold), Paragraph("Groundwater", body_bold)]
    bh_table_data = [bh_headers]
    for b in bhs:
        bh_table_data.append([
            Paragraph(str(b.get("borehole_id")), body_style),
            Paragraph(f"{b.get('cwr_depth')} m", body_style),
            Paragraph(f"{b.get('hard_rock_depth')} m", body_style),
            Paragraph(f"{b.get('termination_depth')} m", body_style),
            Paragraph(str(b.get("groundwater_depth")), body_style)
        ])
    t_bh = Table(bh_table_data, colWidths=[100, 110, 110, 110, 110])
    t_bh.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor("#0284C7")),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor("#E2E8F0")),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor("#F0F9FF")]),
        ('TOPPADDING', (0, 0), (-1, -1), 4),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
    ]))
    story.append(t_bh)
    story.append(Spacer(1, 12))

    # AI Geotechnical Risk Assessment
    story.append(Paragraph("AI Geotechnical Risk Assessment (Interpretation Layer)", h2_style))
    risks = report_data.get("risks", [])
    risk_headers = [Paragraph("Risk Item", body_bold), Paragraph("Severity", body_bold), Paragraph("Reason / Foundation Impact", body_bold), Paragraph("Recommended Action", body_bold)]
    risk_table_data = [risk_headers]
    for r in risks:
        risk_table_data.append([
            Paragraph(r.get("risk_title", ""), body_style),
            Paragraph(r.get("severity", ""), body_bold),
            Paragraph(r.get("reason", ""), body_style),
            Paragraph(r.get("recommended_action", ""), body_style)
        ])
    t_risk = Table(risk_table_data, colWidths=[130, 60, 170, 180])
    t_risk.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor("#DC2626")),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor("#E2E8F0")),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor("#FFF1F2")]),
        ('TOPPADDING', (0, 0), (-1, -1), 4),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
    ]))
    story.append(t_risk)

    doc.build(story)
    return buf.getvalue()
