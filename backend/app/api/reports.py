from typing import Dict, Any, List
from datetime import date, timedelta
from io import BytesIO
from fastapi import APIRouter, Depends, HTTPException, Response
from sqlalchemy.orm import Session
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from backend.app.database.connection import get_db
from backend.app.models.all_models import (
    Project, Activity, DailyProgress, Resource, MLDelayPrediction, ProjectCompletionForecast, Blocker
)

router = APIRouter(prefix="/api/projects/{project_id}/reports", tags=["Reports"])

@router.get("/daily")
def get_daily_report(project_id: int, report_date: str = None, db: Session = Depends(get_db)):
    """Daily Progress & Site Performance Report."""
    r_date = date.fromisoformat(report_date) if report_date else date.today()
    progs = (
        db.query(DailyProgress)
        .filter(DailyProgress.project_id == project_id, DailyProgress.report_date == r_date)
        .all()
    )
    act_map = {a.id: a.name for a in db.query(Activity).filter(Activity.project_id == project_id).all()}

    return {
        "report_type": "Daily Progress Report",
        "date": r_date.isoformat(),
        "entries_count": len(progs),
        "total_workers_assigned": sum(p.workers_assigned for p in progs),
        "total_labour_shortage": sum(p.labour_shortage for p in progs),
        "items": [
            {
                "activity_name": act_map.get(p.activity_id, "Activity"),
                "planned_qty": p.planned_quantity,
                "actual_qty": p.actual_quantity,
                "productivity": p.daily_productivity,
                "variance_pct": p.progress_variance_percent,
                "issues": p.issues,
                "weather": p.weather
            }
            for p in progs
        ]
    }

@router.get("/delay")
def get_delay_report(project_id: int, db: Session = Depends(get_db)):
    """Comprehensive Delay & Slippage Report."""
    delayed = (
        db.query(Activity)
        .filter(Activity.project_id == project_id, Activity.status.in_(["DELAYED", "BLOCKED"]))
        .all()
    )
    preds = {
        p.activity_id: p
        for p in db.query(MLDelayPrediction).filter(MLDelayPrediction.project_id == project_id).all()
    }
    blockers = db.query(Blocker).filter(Blocker.project_id == project_id, Blocker.status != "RESOLVED").all()

    return {
        "report_type": "Delay & Slippage Report",
        "generated_at": date.today().isoformat(),
        "delayed_activities_count": len(delayed),
        "open_blockers_count": len(blockers),
        "activities": [
            {
                "id": a.id,
                "name": a.name,
                "status": a.status,
                "is_critical": a.is_critical,
                "total_float": a.total_float,
                "predicted_delay_days": preds[a.id].predicted_delay_days if a.id in preds else 0.0,
                "delay_risk": preds[a.id].risk_level if a.id in preds else "LOW"
            }
            for a in delayed
        ]
    }

@router.get("/export/excel")
def export_project_excel(project_id: int, db: Session = Depends(get_db)):
    """Generates and downloads full multi-tab executive Excel spreadsheet."""
    project = db.query(Project).filter(Project.id == project_id).first()
    if not project:
        raise HTTPException(status_code=404, detail="Project not found")

    activities = db.query(Activity).filter(Activity.project_id == project_id).order_by(Activity.floor, Activity.sort_order).all()
    resources = db.query(Resource).filter(Resource.project_id == project_id).all()
    delays = {p.activity_id: p for p in db.query(MLDelayPrediction).filter(MLDelayPrediction.project_id == project_id).all()}

    wb = openpyxl.Workbook()

    # Style templates
    header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")

    # Sheet 1: Project Overview
    ws_meta = wb.active
    ws_meta.title = "Overview"
    ws_meta.append(["CONSTRUCTION INTELLIGENCE PLATFORM - PROJECT REPORT"])
    ws_meta.append(["Project Name", project.name])
    ws_meta.append(["Project Code", project.code])
    ws_meta.append(["Construction Type", project.construction_type])
    ws_meta.append(["Location", project.location])
    ws_meta.append(["Total Floors", project.num_floors])
    ws_meta.append(["Built-up Area", f"{project.built_up_area} sq.m"])
    ws_meta.append(["Planned Start", str(project.planned_start_date)])
    ws_meta.append(["Target Completion", str(project.target_completion_date)])
    ws_meta.append(["Project Manager", project.project_manager or "N/A"])

    # Sheet 2: Activities Schedule
    ws_acts = wb.create_sheet(title="Activities & Schedule")
    act_headers = ["ID", "Code", "Activity Name", "Floor", "Tower", "Phase", "Start Date", "End Date", "Dur (days)", "Progress %", "Status", "Critical?", "Delay Risk", "Pred Slip (days)"]
    ws_acts.append(act_headers)
    for col_idx in range(1, len(act_headers) + 1):
        cell = ws_acts.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font

    for a in activities:
        d = delays.get(a.id)
        ws_acts.append([
            a.id, a.code, a.name, a.floor, a.tower, a.phase,
            str(a.start_date), str(a.end_date), a.planned_duration,
            f"{a.progress_percent}%", a.status, "YES" if a.is_critical else "NO",
            d.risk_level if d else "LOW", d.predicted_delay_days if d else 0.0
        ])

    # Sheet 3: Resources
    ws_res = wb.create_sheet(title="Resources")
    res_headers = ["ID", "Category", "Resource Name", "Unit", "Standard Rate ($)", "Available Capacity"]
    ws_res.append(res_headers)
    for col_idx in range(1, len(res_headers) + 1):
        cell = ws_res.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font

    for r in resources:
        ws_res.append([r.id, r.category, r.type_name, r.unit, r.standard_rate, r.available_capacity])

    stream = BytesIO()
    wb.save(stream)
    stream.seek(0)

    filename = f"{project.code}_Construction_Intelligence_Report.xlsx"
    headers = {"Content-Disposition": f'attachment; filename="{filename}"'}
    return Response(
        content=stream.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers=headers
    )
