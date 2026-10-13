"""Schedule import from the 41 ZOY Excel template: validation, all-or-nothing writes, baseline protection, permissions."""
import io
import uuid
from datetime import date, timedelta
from pathlib import Path

import openpyxl
import pytest
from fastapi.testclient import TestClient

from backend.app.database.connection import SessionLocal
from backend.app.main import app
from backend.app.models.all_models import Activity, ActivityDependency, Project
from tests.conftest import auth_headers, grant_membership

SAMPLE = Path(__file__).resolve().parent.parent / "data" / "zoy" / "41_ZOY_schedule_sample.xlsx"
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
client = TestClient(app)
PM, ENG, ADMIN, SM = ("pm@construction.ai", "eng@construction.ai", "admin@construction.ai", "sm@construction.ai")


@pytest.fixture()
def project():
    """A fresh empty project the PM and engineer belong to."""
    db = SessionLocal()
    code = f"IMP-{uuid.uuid4().hex[:6]}"
    p = Project(
        name="Import Test", code=code, construction_type="Residential", location="Pune", built_up_area=1.0,
        planned_start_date=date(2025, 7, 1), target_completion_date=date(2029, 1, 31),
    )
    db.add(p)
    db.commit()
    db.refresh(p)
    grant_membership(db, p.id, PM, "ALL")
    grant_membership(db, p.id, ENG, "Tower A")
    pid = p.id
    db.close()
    return pid


def _sample_bytes(mutate=None) -> bytes:
    wb = openpyxl.load_workbook(SAMPLE)
    if mutate:
        mutate(wb["Schedule"])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _post(pid, content, email=PM, name="zoy.xlsx", **form):
    data = {k: str(v).lower() if isinstance(v, bool) else str(v) for k, v in form.items()}
    return client.post(f"/api/projects/{pid}/schedule/import", files={"file": (name, content, XLSX)}, data=data,
                       headers=auth_headers(email))


def _counts(pid):
    db = SessionLocal()
    try:
        return (db.query(Activity).filter(Activity.project_id == pid).count(),
                db.query(ActivityDependency).filter(ActivityDependency.project_id == pid).count())
    finally:
        db.close()


def _find_row(ws, row_id):
    for row in ws.iter_rows(min_row=2):
        if row[0].value == row_id:
            return row
    raise AssertionError(row_id)


def _expected_from_file():
    ws = openpyxl.load_workbook(SAMPLE)["Schedule"]
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    work = [r for r in rows if r[5] != "Summary"]
    deps = sum(len([t for t in (r[14] or "").split(",") if t]) for r in work)
    return rows, len(work), deps


def test_dry_run_reports_without_writing(project):
    _rows, work, deps = _expected_from_file()
    res = _post(project, _sample_bytes(), dry_run=True)
    assert res.status_code == 200, res.text
    body = res.json()
    assert body["ok"] and body["dry_run"]
    assert body["summary"]["work_rows"] == work
    assert body["summary"]["dependencies"] == deps
    assert set(body["summary"]["activities_per_tower"]) == {"Common", "Tower A", "Tower B", "External / Infra"}
    assert _counts(project) == (0, 0)


def test_import_creates_activities_baseline_and_dependencies(project):
    rows, work, deps = _expected_from_file()
    res = _post(project, _sample_bytes())
    assert res.status_code == 200, res.text
    assert _counts(project) == (work, deps)
    assert res.json()["result"]["schedule"].get("critical_activities_count", 0) > 0

    db = SessionLocal()
    try:
        by_code = {a.code: a for a in db.query(Activity).filter(Activity.project_id == project).all()}
        excavation = next(r for r in rows if r[0] == 5)  # A-Wing EXCAVATION, already completed
        a = by_code[next(c for c in by_code if c.endswith("-005"))]
        assert a.name == "EXCAVATION"
        assert a.tower == "Tower A"
        assert a.start_date == excavation[9].date() and a.end_date == excavation[10].date()
        assert a.baseline_start_date == excavation[7].date() and a.baseline_end_date == excavation[8].date()
        assert a.actual_start_date == excavation[11].date()
        assert a.progress_percent == 100 and a.status == "COMPLETED"
        assert a.baseline_source_file == "zoy.xlsx"
        assert a.phase == "Substructure"
        assert {x.tower for x in by_code.values()} == {"Common", "Tower A", "Tower B", "External / Infra"}
        slab5 = next(a for a in by_code.values() if a.name.startswith("10 SLAB (5TH FLOOR") and a.tower == "Tower A")
        assert slab5.floor == 5
    finally:
        db.close()


def test_reimport_protects_the_baseline(project):
    assert _post(project, _sample_bytes()).status_code == 200
    before = _counts(project)

    refused = _post(project, _sample_bytes())
    assert refused.status_code == 409 and "baseline" in refused.json()["detail"].lower()
    assert _counts(project) == before

    # a deliberate new baseline updates in place, it does not duplicate
    assert _post(project, _sample_bytes(), new_baseline=True).status_code == 200
    assert _counts(project) == before


def test_other_activities_block_import_unless_admin_replaces(project):
    db = SessionLocal()
    db.add(Activity(project_id=project, name="Old plan item", code="OLD-1", phase="Execution", tower="Tower A",
                    start_date=date.today(), end_date=date.today() + timedelta(days=1)))
    db.commit()
    db.close()

    assert _post(project, _sample_bytes()).status_code == 409
    assert _post(project, _sample_bytes(), replace_existing=True).status_code == 403  # PM may not
    ok = _post(project, _sample_bytes(), email=ADMIN, replace_existing=True)
    assert ok.status_code == 200, ok.text
    assert ok.json()["result"]["replaced_existing"] == 1
    db = SessionLocal()
    assert db.query(Activity).filter(Activity.project_id == project, Activity.code == "OLD-1").count() == 0
    db.close()


def _set(row_id, column, value):
    def mutate(ws):
        headers = [c.value for c in ws[1]]
        _find_row(ws, row_id)[headers.index(column)].value = value
    return mutate


def _drop_column(name):
    def mutate(ws):
        ws.cell(row=1, column=[c.value for c in ws[1]].index(name) + 1).value = "Renamed"
    return mutate


@pytest.mark.parametrize(
    "mutate,field_name,fragment",
    [
        (_drop_column("Predecessors"), "header", "Missing column"),
        (_set(10, "ID", 9), "ID", "Duplicate ID"),
        (_set(10, "Predecessors", "9999"), "Predecessors", "does not exist"),
        (_set(10, "Predecessors", "10"), "Predecessors", "depend on itself"),
        (_set(10, "Predecessors", "3"), "Predecessors", "summary row"),
        (_set(10, "Predecessors", "banana"), "Predecessors", "cannot read predecessor"),
        (_set(9, "Predecessors", "10"), "Predecessors", "Circular dependency"),
        (_set(10, "Wing / Area", "C-Wing"), "Wing / Area", "Unknown wing"),
        (_set(10, "Finish", date(2020, 1, 1)), "Finish", "before Start"),
        (_set(10, "% Complete", 1.5), "% Complete", "between 0% and 100%"),
        (_set(10, "Duration (days)", "many"), "Duration (days)", "number of days"),
        (_set(10, "Start", "not a date"), "Start", "not a date"),
        (_set(10, "Task Name", ""), "Task Name", "empty"),
    ],
    ids=["missing-column", "duplicate-id", "unknown-pred", "self-pred", "summary-pred", "bad-pred-syntax",
         "cycle", "unknown-wing", "finish-before-start", "percent-range", "bad-duration", "bad-date", "empty-name"],
)
def test_invalid_files_are_rejected_and_write_nothing(project, mutate, field_name, fragment):
    res = _post(project, _sample_bytes(mutate))
    assert res.status_code == 422, res.text
    errors = res.json()["errors"]
    assert any(e["field"] == field_name and fragment.lower() in e["message"].lower() for e in errors), errors
    assert _counts(project) == (0, 0)


def test_all_errors_are_listed_together(project):
    def mutate(ws):
        _set(10, "Wing / Area", "C-Wing")(ws)
        _set(11, "% Complete", 7)(ws)
    errors = _post(project, _sample_bytes(mutate)).json()["errors"]
    assert {e["row"] for e in errors} >= {_row_number(10), _row_number(11)}


def _row_number(row_id):
    ws = openpyxl.load_workbook(SAMPLE)["Schedule"]
    return _find_row(ws, row_id)[0].row


def test_not_an_excel_file(project):
    assert _post(project, b"id,name\n1,x\n", name="schedule.csv").status_code == 400
    res = _post(project, b"this is not a zip", name="schedule.xlsx")
    assert res.status_code == 422 and res.json()["errors"][0]["field"] == "file"


def test_permissions(project):
    content = _sample_bytes()
    assert _post(project, content, email=ENG).status_code == 403   # engineer role
    assert _post(project, content, email=SM).status_code == 403    # not allowed to import
    assert _counts(project) == (0, 0)

    db = SessionLocal()
    other = Project(name="Other", code=f"OTH-{uuid.uuid4().hex[:6]}", construction_type="Residential", location="x",
                    built_up_area=1.0, planned_start_date=date.today(), target_completion_date=date.today())
    db.add(other)
    db.commit()
    other_id = other.id
    db.close()
    assert _post(other_id, content, email=PM).status_code == 403   # PM is not a member of that project
    assert _post(project, content, email=PM, new_baseline=True).status_code == 200  # PM may confirm a baseline
