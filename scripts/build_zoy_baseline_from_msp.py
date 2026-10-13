"""Build the 41 ZOY baseline workbook from the MSP master schedule tracker (06-Oct-2026 PDF).

Real data (from the PDF): MSP IDs, outline levels, durations, start / finish dates, predecessors with lags.
Names, % complete and which rows have actual dates come from the earlier tracker sheet (scripts/generate_zoy_sample_schedule.py
RAW table); both lists have the same 154 rows in the same order, which is checked below.

The PDF export hides child rows of collapsed tasks (e.g. the floor-wise blockwork). Predecessors that point at a hidden
row are re-pointed to the nearest visible task above it, as start-to-start (or finish-to-finish) with the lag recalculated
from the visible dates. Every such change is listed on the Read Me sheet.

Usage: python scripts/build_zoy_baseline_from_msp.py [output.xlsx]
"""
import re
import sys
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

sys.path.insert(0, str(Path(__file__).resolve().parent))
import generate_zoy_sample_schedule as old  # noqa: E402  (names, % complete and actual flags)

TODAY = date(2026, 10, 12)

# id | outline level | duration | start | finish | predecessors            (MSP export, "Baseline Master Schedule Tracker 06-OCT-2026")
PDF = """
1|1|1066.5|7/01/25|1/18/29|
2|2|75|7/01/25|9/30/25|
10|2|1057|7/01/25|1/06/29|
11|3|1|7/01/25|7/01/25|
12|3|113|7/02/25|11/19/25|11
13|3|728|11/20/25|4/20/28|
14|4|104|11/20/25|3/25/26|
24|4|596|2/25/26|2/19/28|
25|5|41|2/25/26|4/16/26|23SS+8 days
26|5|84|4/06/26|7/13/26|
33|5|163.8|4/18/26|11/02/26|
40|5|16.2|7/02/26|8/14/26|32FS-5 days
41|5|18|8/09/26|10/10/26|40FS-5 days
42|5|20|9/23/26|10/30/26|41FS-5 days
43|5|4|5/25/26|5/28/26|44FS-10 days
44|5|0|6/18/26|6/18/26|45FS-70 days
45|5|1|10/13/26|10/13/26|46SS-10 days
46|5|35|10/26/26|12/08/26|42FS-5 days
47|5|20|10/31/26|11/26/26|42
48|5|25|11/27/26|12/25/26|47
49|5|20|12/26/26|1/18/27|48
50|5|13|1/19/27|2/03/27|49
51|5|13|2/04/27|2/18/27|50
52|5|12|2/19/27|3/04/27|51
53|5|12|3/05/27|3/18/27|52
54|5|12|3/19/27|4/03/27|53
55|5|12|4/05/27|4/19/27|54
56|5|12|4/20/27|5/04/27|55
57|5|12|5/05/27|5/18/27|56
58|5|12|5/19/27|6/01/27|57
59|5|12|6/02/27|6/15/27|58
60|5|12|6/16/27|6/29/27|59
61|5|12|6/30/27|7/13/27|60
62|5|14|7/14/27|7/29/27|61
63|5|18|7/30/27|8/19/27|62
64|5|14|8/20/27|9/06/27|63
65|5|12|9/07/27|9/22/27|64
66|5|12|9/23/27|10/06/27|65
67|5|12|10/07/27|10/21/27|66
68|5|12|10/22/27|11/09/27|67
69|5|12|11/10/27|11/23/27|68
70|5|12|11/24/27|12/07/27|69
71|5|12|12/08/27|12/21/27|70
72|5|12|12/22/27|1/04/28|71
73|5|12|1/05/28|1/18/28|72
74|5|12|1/19/28|2/02/28|73
75|5|15|2/03/28|2/19/28|74
76|4|50|2/21/28|4/20/28|
81|3|402|1/19/27|5/20/28|
119|3|349|2/24/27|4/22/28|
155|3|318|6/07/27|6/24/28|
167|3|808|12/08/25|8/10/28|
207|3|471|2/25/27|9/19/28|
245|3|461|4/13/27|10/25/28|
325|3|398|7/20/27|11/17/28|
366|3|493|4/20/27|12/08/28|
372|3|369|8/11/27|11/06/28|
377|3|846|2/25/26|12/20/28|
459|3|387|8/17/27|12/02/28|
505|3|60|7/29/28|10/12/28|
508|3|45|8/26/28|10/24/28|461FS-20 days;556SS+15 days
509|3|35|11/23/28|1/02/29|556FS+5 days
510|3|25|12/04/28|1/01/29|502
511|3|484|4/13/27|11/21/28|
550|3|341|7/14/27|8/30/28|
555|3|60|9/20/28|12/02/28|207
556|3|120|6/21/28|11/16/28|497FF+15 days
557|3|100|7/14/28|11/16/28|556SS+20 days
558|3|45|10/19/28|12/09/28|557FF+20 days
559|3|60|9/28/28|12/09/28|557FF+20 days
560|3|165|5/22/28|12/08/28|
563|3|15|12/21/28|1/06/29|
574|3|60|10/30/28|1/06/29|
578|2|1050.6|7/18/25|1/17/29|
579|3|1|7/18/25|7/18/25|11FS+15 days
580|3|90|7/19/25|11/10/25|579
581|3|753.6|11/11/25|5/12/28|
582|4|88|11/11/25|2/21/26|
592|4|624.6|2/12/26|3/13/28|
593|5|1|2/12/26|4/14/26|591SS+8 days
594|5|61|3/20/26|5/30/26|
601|5|19.5|5/30/26|8/03/26|594
602|5|18|8/03/26|11/03/26|601FS-3 days
603|5|20|9/28/26|11/24/26|602FS-2 days
604|5|4|5/25/26|5/28/26|605FS-10 days
605|5|0|6/18/26|6/18/26|606FS-70 days
606|5|1|11/03/26|11/04/26|607SS-10 days
607|5|35|11/18/26|12/29/26|603FS-5 days
608|5|20|11/24/26|12/17/26|603
609|5|25|12/17/26|1/15/27|608
610|5|20|1/15/27|2/09/27|609
611|5|13|2/09/27|2/24/27|610
612|5|13|2/24/27|3/11/27|611
613|5|12|3/11/27|3/27/27|612
614|5|12|3/27/27|4/12/27|613
615|5|12|4/12/27|4/26/27|614
616|5|12|4/26/27|5/11/27|615
617|5|12|5/11/27|5/25/27|616
618|5|12|5/25/27|6/08/27|617
619|5|12|6/08/27|6/22/27|618
620|5|12|6/22/27|7/06/27|619
621|5|12|7/06/27|7/20/27|620
622|5|12|7/20/27|8/03/27|621
623|5|14|8/03/27|8/19/27|622
624|5|18|8/19/27|9/10/27|623
625|5|14|9/10/27|9/29/27|624
626|5|12|9/29/27|10/14/27|625
627|5|12|10/14/27|11/02/27|626
628|5|12|11/02/27|11/16/27|627
629|5|12|11/16/27|11/30/27|628
630|5|12|11/30/27|12/14/27|629
631|5|12|12/14/27|12/28/27|630
632|5|12|12/28/27|1/11/28|631
633|5|12|1/11/28|1/25/28|632
634|5|12|1/25/28|2/09/28|633
635|5|12|2/09/28|2/23/28|634
636|5|15|2/23/28|3/13/28|635
637|4|50|3/13/28|5/12/28|
642|3|402|2/09/27|6/10/28|
680|3|349|3/16/27|5/15/28|
716|3|318|6/26/27|7/15/28|
728|3|833.6|11/28/25|9/04/28|
768|3|471|3/17/27|10/11/28|
806|3|461|5/04/27|11/15/28|
886|3|398|8/09/27|12/08/28|
927|3|493|5/11/27|12/29/28|
933|3|369|8/31/27|11/27/28|
938|3|857.6|3/04/26|1/10/29|
1020|3|387|9/07/27|12/23/28|
1066|3|60|8/19/28|11/06/28|
1069|3|45|9/18/28|11/14/28|1022FS-20 days;1117SS+15 days
1070|3|45|11/25/28|1/17/29|1117FS-10 days
1071|3|30|11/30/28|1/04/29|1063FF+10 days
1072|3|484|5/04/27|12/12/28|
1111|3|341|8/03/27|9/22/28|
1116|3|60|10/11/28|12/23/28|768
1117|3|120|7/11/28|12/07/28|1058FF+15 days
1118|3|100|8/03/28|12/07/28|1117SS+20 days
1119|3|40|11/14/28|12/30/28|1118FF+20 days
1120|3|60|10/21/28|12/30/28|1118FF+20 days
1121|3|165|6/10/28|12/29/28|
1124|3|15|12/23/28|1/10/29|
1135|3|44|11/18/28|1/09/29|
1139|2|861.5|3/10/26|1/18/29|
1140|3|681.5|3/10/26|6/13/28|
1141|4|90|3/10/26|11/17/26|25FS+5 days
1142|4|100|1/14/28|5/15/28|1141FS+350 days
1143|4|25|5/15/28|6/13/28|1142
1144|3|545|11/27/26|9/15/28|
1148|3|250|9/10/27|7/10/28|1145FS+90 days
1149|3|180|4/24/28|11/29/28|119
1150|3|200|5/09/28|1/06/29|1143FS-30 days
1151|3|210|5/09/28|1/18/29|1143FS-30 days
1152|3|120|7/11/28|12/06/28|1148
"""

WING_OF_TOP = {"A-WING": "A-Wing", "B-WING": "B-Wing", "EXTERNAL": "External Development"}


def d(text):
    return datetime.strptime(text, "%m/%d/%y").date()


class R:
    pass


def load():
    old_rows = [old.Row(line) for line in old.RAW.strip().splitlines()]
    rows = []
    for k, line in enumerate(PDF.strip().splitlines()):
        pid, outline, dur, start, finish, preds = line.split("|")
        o = old_rows[k]
        r = R()
        r.id, r.level, r.dur = int(pid), int(outline) - 1, float(dur)
        r.name, r.start, r.finish, r.preds_text = o.name, d(start), d(finish), preds
        r.pct, r.has_as, r.has_af = o.pct, o.as_m is not None, o.af_m is not None
        assert abs(o.dur - r.dur) < 0.01, f"row {k}: duration {o.dur} vs {r.dur} for {o.name}"
        rows.append(r)
    for i, r in enumerate(rows):
        nxt = rows[i + 1] if i + 1 < len(rows) else None
        r.summary = bool(nxt and nxt.level > r.level)
        r.kind = "Summary" if r.summary else ("Milestone" if r.dur == 0 else "Task")
    wing = "Project"
    for r in rows:
        if r.level == 1:
            wing = next((v for k, v in WING_OF_TOP.items() if r.name.upper().startswith(k)), "Project")
        elif r.level == 0:
            wing = "Project"
        r.wing = wing
    return rows


def work_days(a, b):
    """Calendar days between two dates converted to working days (6-day week), rounded."""
    return round((b - a).days * 6 / 7)


TOKEN = re.compile(r"^(\d+)\s*(FS|SS|FF|SF)?\s*(?:([+-])\s*(\d+)\s*days?)?$", re.I)


def resolve_predecessors(rows):
    by_id = {r.id: r for r in rows}
    ids = sorted(by_id)
    notes = []
    for r in rows:
        out = []
        for token in [t.strip() for t in r.preds_text.split(";") if t.strip()]:
            m = TOKEN.match(token)
            assert m, f"unreadable predecessor {token!r} on {r.id}"
            pid, kind = int(m.group(1)), (m.group(2) or "FS").upper()
            lag = int(m.group(4) or 0) * (-1 if m.group(3) == "-" else 1)
            if pid in by_id:
                target = by_id[pid]
                if target.summary:
                    notes.append(f"{r.id} {r.name}: predecessor {pid} ({target.name}) is a summary row, link dropped")
                    continue
                out.append((pid, kind, lag))
                continue
            # hidden row: the visible task just above it owns it
            parent = by_id[max(i for i in ids if i < pid)]
            if parent.summary or parent.id == r.id:
                notes.append(f"{r.id} {r.name}: predecessor {pid} is hidden and its parent {parent.id} cannot be used, link dropped")
                continue
            new_kind = "FF" if kind == "FF" else "SS"
            if new_kind == "FF":
                new_lag = work_days(parent.finish, r.finish)
            else:
                new_lag = work_days(parent.start, r.start)
            out.append((parent.id, new_kind, new_lag))
            notes.append(
                f"{r.id} {r.name}: predecessor {pid}{kind}{lag:+d}d is a hidden sub-task of {parent.id} ({parent.name}); "
                f"linked to {parent.id} as {new_kind}{new_lag:+d}d"
            )
        r.preds = out
    return notes


def fmt_pred(r):
    parts = []
    for pid, kind, lag in r.preds:
        t = f"{pid}" if kind == "FS" and lag == 0 else f"{pid}{kind}"
        if lag:
            t += f"{lag:+d}d"
        parts.append(t)
    return ",".join(parts)


def actuals(r):
    """Sample progress as at TODAY: actual dates follow the plan for rows the tracker shows as started / finished."""
    a_start = r.start if r.has_as and r.start <= TODAY else None
    a_finish = min(r.finish, TODAY) if r.has_af and a_start else None
    return a_start, a_finish


SITE_ENGINEER = {"A-Wing": "Aditya Patil (Sr. Engineer)", "B-Wing": "Vicky Satpute (Sr. Engineer)", "External Development": "To be appointed (Jan-27)"}


def write(rows, notes, path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Schedule"
    ws.append(["ID", "WBS", "Level", "Task Name", "Wing / Area", "Type", "Duration (days)", "Baseline Start",
               "Baseline Finish", "Start", "Finish", "Actual Start", "Actual Finish", "% Complete", "Predecessors", "Site Engineer"])
    counters = [0] * 6
    for r in rows:
        counters[r.level] += 1
        for k in range(r.level + 1, 6):
            counters[k] = 0
        wbs = ".".join(str(c) for c in counters[1:r.level + 1]) if r.level else "0"
        a_start, a_finish = actuals(r)
        ws.append([r.id, wbs, r.level, ("    " * r.level) + r.name, r.wing, r.kind, r.dur, r.start, r.finish, r.start, r.finish,
                   a_start, a_finish, r.pct / 100, fmt_pred(r), SITE_ENGINEER.get(r.wing, "")])
    fill = PatternFill("solid", fgColor="1F3A5F")
    for c in ws[1]:
        c.font, c.fill = Font(bold=True, color="FFFFFF"), fill
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    summary_ids = {r.id for r in rows if r.summary}
    for row in ws.iter_rows(min_row=2):
        for c in row[7:13]:
            c.number_format = "dd-mmm-yyyy"
        row[13].number_format = "0%"
        if row[0].value in summary_ids:
            for c in row:
                c.font = Font(bold=True)
                c.fill = PatternFill("solid", fgColor="E8EEF5")
        ws.row_dimensions[row[0].row].outlineLevel = max(row[2].value - 1, 0)
    for i, w in enumerate([7, 8, 7, 62, 20, 11, 14, 15, 15, 15, 15, 15, 15, 12, 26, 28], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "E2"
    ws.auto_filter.ref = ws.dimensions

    readme = wb.create_sheet("Read Me", 0)
    lines = [
        ("41 ZOY - Baseline master schedule (from MSP tracker 06-Oct-2026)", True),
        ("", False),
        ("Source: ZOY - Baseline Master Schedule Tracker 06-OCT-2026.pdf. IDs are the MSP task IDs, so predecessors use them directly.", False),
        ("Dates, durations and predecessors are real. Start / Finish are also used as the baseline (the tracker holds one set of dates).", False),
        (f"Progress and actual dates are SAMPLE data as at {TODAY:%d-%b-%Y}, taken from the earlier tracker sheet (which tasks were started / finished and their % complete).", False),
        ("Predecessors: '12' = finish-to-start, '23SS+8d' = start-to-start with 8 days lag, '497FF+15d' = finish-to-finish, '+/-' lags are in working days.", False),
        ("", False),
        ("Hidden rows: the PDF hides the children of collapsed tasks (for example the floor-wise blockwork). Links to hidden rows were re-pointed:", True),
    ]
    for n in notes:
        lines.append(("- " + n, False))
    lines += [("", False), ("To get the full plan, export the MSP file with all outline levels expanded (or as XML) and import that instead.", True)]
    for text, bold in lines:
        readme.append([text])
        readme.cell(readme.max_row, 1).font = Font(bold=bold, size=13 if readme.max_row == 1 else 11)
    readme.column_dimensions["A"].width = 150
    wb.active = 1
    wb.save(path)


def main():
    out = Path(sys.argv[1] if len(sys.argv) > 1 else "data/zoy/41_ZOY_baseline_master_2026-10-06.xlsx")
    out.parent.mkdir(parents=True, exist_ok=True)
    rows = load()
    notes = resolve_predecessors(rows)
    write(rows, notes, out)
    deps = sum(len(r.preds) for r in rows)
    print(f"wrote {out}: {len(rows)} rows, {sum(1 for r in rows if not r.summary)} work rows, {deps} dependencies")
    for n in notes:
        print(" ", n)


if __name__ == "__main__":
    main()
