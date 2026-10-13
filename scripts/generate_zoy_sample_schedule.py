"""Generate the sample 41 ZOY schedule workbook (real dates + predecessors).

The MSP export shared by planning had masked dates ("D/7/YYYY": only the month survived).
This script rebuilds a coherent best-case schedule from what the export does contain:
task names, outline levels, durations (6-day work week), month hints and % complete.
The result is the template future uploads should follow.

Usage: python scripts/generate_zoy_sample_schedule.py [output.xlsx]
"""
import sys
from datetime import date, timedelta
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

TODAY = date(2026, 10, 12)
PROJECT_START = date(2025, 7, 1)
PROJECT_FINISH_LIMIT = date(2029, 1, 31)  # current finish month of the project row is January

# sr | level | name | baseline start month | baseline finish month | duration (working days)
#    | start month | finish month | actual start month | actual finish month | % complete   ("-" = NA)
RAW = """
1|0|41 ZOY|7|11|1066.5|7|1|7|-|6
2|1|SITE MOBILIZATION|7|9|75|7|9|7|9|100
3|1|A-WING (B+Stilt+P1+P2+P3+P4+29F+Terrace)|7|11|1057|7|1|7|-|6
4|2|START OF CONSTRUCTION WORKS|7|7|1|7|7|7|7|100
5|2|EXCAVATION|7|11|113|7|11|7|11|100
6|2|RCC WORK|11|3|728|11|4|11|-|41
7|3|SUB- STRUCTURE|11|2|104|11|3|11|3|100
8|3|SUPER- STRUCTURE|2|3|596|2|2|2|-|31
9|4|01 SLAB ( STILT FLOOR SLAB)|2|4|41|2|4|2|4|100
10|4|02 SLAB ( PODIUM 1 FLOOR SLAB)|4|5|84|4|7|4|7|100
11|4|EXTENDED BASEMENT|4|6|163.8|4|11|4|-|47
12|4|03 SLAB ( PODIUM 2 FLOOR SLAB)|5|6|16.2|7|8|7|8|100
13|4|04 SLAB ( PODIUM 3 FLOOR SLAB)|6|7|18|8|10|8|-|80
14|4|05 SLAB ( 1ST FLOOR)|7|8|20|9|10|9|-|0
15|4|ADVANCE FOR ALUFORM|5|5|4|5|5|5|5|100
16|4|SHELL PLAN FINALIZATION|2|2|0|6|6|6|6|100
17|4|ALUFORM MATERIAL DELIVERY|7|7|1|10|10|-|-|0
18|4|06 SLAB (2ND FLOOR ) - PART ALUFORM MOCK UP|8|10|35|10|12|-|-|0
19|4|06 SLAB (2ND FLOOR ) - PART CONVENTIONAL|8|9|20|10|11|-|-|0
20|4|07 SLAB (3RD FLOOR ) - COMPLETE ALUFORM MOCK UP|9|10|25|11|12|-|-|0
21|4|09 SLAB (4TH FLOOR ) REFF.|10|11|20|12|1|-|-|0
22|4|10 SLAB (5TH FLOOR )|11|12|13|1|2|-|-|0
23|4|11 SLAB (6TH FLOOR)|12|12|13|2|2|-|-|0
24|4|12 SLAB (7TH FLOOR )|12|1|12|2|3|-|-|0
25|4|13 SLAB (8TH FLOOR )|1|1|12|3|3|-|-|0
26|4|14 SLAB (9TH FLOOR ) REFF.|1|2|12|3|4|-|-|0
27|4|15 SLAB (10TH FLOOR)|2|2|12|4|4|-|-|0
28|4|16 SLAB (11TH FLOOR )|2|3|12|4|5|-|-|0
29|4|17 SLAB (12TH FLOOR)|3|3|12|5|5|-|-|0
30|4|18 SLAB (13TH FLOOR)|3|4|12|5|6|-|-|0
31|4|19 SLAB (14TH FLOOR) REFF.|4|4|12|6|6|-|-|0
32|4|20 SLAB (15TH FLOOR)|4|5|12|6|6|-|-|0
33|4|21 SLAB (16TH FLOOR)|5|5|12|6|7|-|-|0
34|4|22 SLAB (17TH FLOOR)|5|6|14|7|7|-|-|0
35|4|23 SLAB (18TH FLOOR) (RCC Water Tank)|6|7|18|7|8|-|-|0
36|4|24 SLAB (19TH FLOOR) REFF.|7|7|14|8|9|-|-|0
37|4|25 SLAB (20TH FLOOR)|7|8|12|9|9|-|-|0
38|4|26 SLAB (21ST FLOOR)|8|8|12|9|10|-|-|0
39|4|27 SLAB (22ND FLOOR)|8|9|12|10|10|-|-|0
40|4|28 SLAB(23RD FLOOR)|9|9|12|10|11|-|-|0
41|4|29 SLAB(24TH FLOOR) REFF.|9|10|12|11|11|-|-|0
42|4|30 SLAB(25TH FLOOR)|10|11|12|11|12|-|-|0
43|4|31 SLAB(26TH FLOOR)|11|11|12|12|12|-|-|0
44|4|32 SLAB(27TH FLOOR)|11|12|12|12|1|-|-|0
45|4|33 SLAB(28TH FLOOR)|12|12|12|1|1|-|-|0
46|4|34 SLAB(29TH FLOOR)|12|1|12|1|2|-|-|0
47|4|35 SLAB (TERRACE FLOOR)|1|1|15|2|2|-|-|0
48|3|ABOVE TERRACE RCC|1|3|50|2|4|-|-|0
49|2|BLOCKWORK|11|4|402|1|5|-|-|0
50|2|INTERNAL PLASTER (GYPSUM)|12|3|349|2|4|-|-|0
51|2|EXTERNAL GRINDING|4|5|318|6|6|-|-|0
52|2|WATERPROOFING|12|7|808|12|8|12|-|7
53|2|PLUMBING WORK|12|8|471|2|9|-|-|0
54|2|TILING WORK|2|9|461|4|10|-|-|0
55|2|CARPENTRY WORK|5|10|398|7|11|-|-|0
56|2|FABRICATION|2|11|493|4|12|-|-|0
57|2|ALUMINIUM WINDOWS|6|10|369|8|11|-|-|0
58|2|ELECTRICAL WORK|4|10|846|2|12|2|-|5
59|2|PAINTING WORK|6|10|387|8|12|-|-|0
60|2|FALSE CEILING|5|8|60|7|10|-|-|0
61|2|SANITARY FIXING|6|8|45|8|10|-|-|0
62|2|CP|9|11|35|11|1|-|-|0
63|2|VDP|10|11|25|12|1|-|-|0
64|2|FIRE FIGHTING|2|9|484|4|11|-|-|0
65|2|LIFT WORK|5|7|341|7|8|-|-|0
66|2|SOLAR|8|11|60|9|12|-|-|0
67|2|ACID WASH|4|9|120|6|11|-|-|0
68|2|TILE GROUTING|5|9|100|7|11|-|-|0
69|2|SNAGGING & DESNAGGING|9|11|45|10|12|-|-|0
70|2|DEEP CLEANING|9|11|60|9|12|-|-|0
71|2|ELEVATION FEATURES|4|11|165|5|12|-|-|0
72|2|OTHER SERVICE|10|11|15|12|1|-|-|0
73|2|NOC'S|9|11|60|10|1|-|-|0
74|1|B-WING (B+Stilt+P1+P2+P3+P4+29F+Terrace)|7|11|1050.6|7|1|7|-|5
75|2|START OF CONSTRUCTION WORKS|7|7|1|7|7|7|7|100
76|2|EXCAVATION|7|11|90|7|11|7|11|100
77|2|RCC WORK|11|3|753.6|11|5|11|-|37
78|3|SUB- STRUCTURE|11|2|88|11|2|11|2|100
79|3|SUPER- STRUCTURE|2|1|624.6|2|3|2|-|17
80|4|01 SLAB ( STILT FLOOR SLAB)|2|4|1|2|4|2|4|100
81|4|02 SLAB ( PODIUM 1 FLOOR SLAB)|4|5|61|3|5|3|5|100
82|4|03 SLAB ( PODIUM 2 FLOOR SLAB)|5|6|19.5|5|8|5|8|100
83|4|04 SLAB ( PODIUM 3 FLOOR SLAB)|6|6|18|8|11|8|-|80
84|4|05 SLAB ( 1ST FLOOR)|7|8|20|9|11|9|-|15
85|4|ADVANCE FOR ALUFORM|-|-|4|5|5|5|5|100
86|4|SHELL PLAN FINALIZATION|2|2|0|6|6|6|6|100
87|4|ALUFORM MATERIAL DELIVERY|7|7|1|11|11|-|-|0
88|4|06 SLAB (2ND FLOOR ) - PART ALUFORM MOCK UP|8|10|35|11|12|-|-|0
89|4|06 SLAB (2ND FLOOR ) - PART CONVENTIONAL|8|9|20|11|12|-|-|0
90|4|07 SLAB (3RD FLOOR ) - COMPLETE ALUFORM MOCK UP|9|10|25|12|1|-|-|0
91|4|09 SLAB (4TH FLOOR ) REFF.|10|11|20|1|2|-|-|0
92|4|10 SLAB (5TH FLOOR )|11|11|13|2|2|-|-|0
93|4|11 SLAB (6TH FLOOR)|11|12|13|2|3|-|-|0
94|4|12 SLAB (7TH FLOOR )|12|12|12|3|3|-|-|0
95|4|13 SLAB (8TH FLOOR )|12|1|12|3|4|-|-|0
96|4|14 SLAB (9TH FLOOR ) REFF.|1|1|12|4|4|-|-|0
97|4|15 SLAB (10TH FLOOR)|1|2|12|4|5|-|-|0
98|4|16 SLAB (11TH FLOOR )|2|2|12|5|5|-|-|0
99|4|17 SLAB (12TH FLOOR)|2|3|12|5|6|-|-|0
100|4|18 SLAB (13TH FLOOR)|3|3|12|6|6|-|-|0
101|4|19 SLAB (14TH FLOOR) REFF.|3|4|12|6|7|-|-|0
102|4|20 SLAB (15TH FLOOR)|4|5|12|7|7|-|-|0
103|4|21 SLAB (16TH FLOOR)|5|5|12|7|8|-|-|0
104|4|22 SLAB (17TH FLOOR)|5|6|14|8|8|-|-|0
105|4|23 SLAB (18TH FLOOR) (RCC Water Tank)|6|6|18|8|9|-|-|0
106|4|24 SLAB (19TH FLOOR) REFF.|6|7|14|9|9|-|-|0
107|4|25 SLAB (20TH FLOOR)|7|7|12|9|10|-|-|0
108|4|26 SLAB (21ST FLOOR)|7|8|12|10|11|-|-|0
109|4|27 SLAB (22ND FLOOR)|8|8|12|11|11|-|-|0
110|4|28 SLAB(23RD FLOOR)|8|9|12|11|11|-|-|0
111|4|29 SLAB(24TH FLOOR) REFF.|9|10|12|11|12|-|-|0
112|4|30 SLAB(25TH FLOOR)|10|10|12|12|12|-|-|0
113|4|31 SLAB(26TH FLOOR)|10|11|12|12|1|-|-|0
114|4|32 SLAB(27TH FLOOR)|11|11|12|1|1|-|-|0
115|4|33 SLAB(28TH FLOOR)|11|12|12|1|2|-|-|0
116|4|34 SLAB(29TH FLOOR)|12|12|12|2|2|-|-|0
117|4|35 SLAB (TERRACE FLOOR)|12|1|15|2|3|-|-|0
118|3|ABOVE TERRACE RCC|1|3|50|3|5|-|-|0
119|2|BLOCKWORK|11|4|402|2|6|-|-|0
120|2|INTERNAL PLASTER (GYPSUM)|12|3|349|3|5|-|-|0
121|2|EXTERNAL GRINDING|4|5|318|6|7|-|-|0
122|2|WATERPROOFING|11|7|833.6|11|9|11|-|8
123|2|PLUMBING WORK|12|8|471|3|10|-|-|0
124|2|TILING WORK|1|9|461|5|11|-|-|0
125|2|CARPENTRY WORK|5|10|398|8|12|-|-|0
126|2|FABRICATION|2|10|493|5|12|-|-|0
127|2|ALUMINIUM WINDOWS|6|9|369|8|11|-|-|0
128|2|ELECTRICAL WORK|4|10|857.6|3|1|3|-|2
129|2|PAINTING WORK|6|9|387|9|12|-|-|0
130|2|FALSE CEILING|5|8|60|8|11|-|-|0
131|2|SANITARY FIXING|6|8|45|9|11|-|-|0
132|2|CP|9|11|45|11|1|-|-|0
133|2|VDP|9|11|30|11|1|-|-|0
134|2|FIRE FIGHTING|1|9|484|5|12|-|-|0
135|2|LIFT WORK|5|7|341|8|9|-|-|0
136|2|SOLAR|8|10|60|10|12|-|-|0
137|2|ACID WASH|4|9|120|7|12|-|-|0
138|2|TILE GROUTING|5|9|100|8|12|-|-|0
139|2|SNAGGING & DESNAGGING|9|11|40|11|12|-|-|0
140|2|DEEP CLEANING|9|11|60|10|12|-|-|0
141|2|ELEVATION FEATURES|4|11|165|6|12|-|-|0
142|2|OTHER SERVICE|10|11|15|12|1|-|-|0
143|2|NOC'S|8|11|44|11|1|-|-|0
144|1|EXTERNAL DEVELOPMENT WORK & SERVICES INCLUDING UTILITIES|4|10|861.5|3|1|3|-|4
145|2|UGWT|4|3|681.5|3|6|3|-|27
146|3|CIVIL WORKS|4|8|90|3|11|3|-|65
147|3|ELECTRO- MECHANICAL WORKS|10|2|100|1|5|-|-|0
148|3|TESTING& COMMISSIONING|2|3|25|5|6|-|-|0
149|2|STP|9|7|545|11|9|-|-|0
150|2|COMPOUND WALL|7|5|250|9|7|-|-|0
151|2|AMENITIES (Civil + Hardscape + Softscape) (Ground+Podium )|3|10|180|4|11|-|-|0
152|2|PLUMBING & SEWERAGE WORKS|1|9|200|5|1|-|-|0
153|2|TRANSFORMER ROOM/ SUB STATION/ EXTERNAL ELECTRICAL WORK|1|10|210|5|1|-|-|0
154|2|ENTRANCE GATE WITH SECURITY CABIN & ENTRY GATE - 2 Nos|5|9|120|7|12|-|-|0
"""

# Chain rows are scheduled one after another (slab cycle). Everything else is placed by its month hints.
CHAIN_A = [9, 10, 12, 13, 14, 17] + list(range(18, 49))
CHAIN_B = [80, 81, 82, 83, 84, 87] + list(range(88, 119))
SITE_ENGINEER = {  # from the 41 ZOY site organisation chart
    "A": "Aditya Patil (Sr. Engineer)",
    "B": "Vicky Satpute (Sr. Engineer)",
    "EXT": "To be appointed (Jan-27)",
}


def month(v):
    return None if v == "-" else int(v)


class Row:
    def __init__(self, line):
        p = line.split("|")
        self.sr, self.level, self.name = int(p[0]), int(p[1]), p[2]
        self.bs_m, self.bf_m = month(p[3]), month(p[4])
        self.dur = float(p[5])
        self.s_m, self.f_m = month(p[6]), month(p[7])
        self.as_m, self.af_m = month(p[8]), month(p[9])
        self.pct = int(p[10])
        self.is_summary = False
        self.block = "ZOY"
        self.start = self.finish = None
        self.b_start = self.b_finish = None
        self.a_start = self.a_finish = None
        self.preds = []  # (sr, type, lag_days)


ROWS = [Row(l) for l in RAW.strip().splitlines()]
BY_SR = {r.sr: r for r in ROWS}


def work_days_after(start, n):
    """Date n working days after start (6-day week, Sunday off)."""
    d = start
    while n > 0:
        d += timedelta(days=1)
        if d.weekday() != 6:
            n -= 1
    return d


def first_workday(d):
    return d + timedelta(days=1) if d.weekday() == 6 else d


def finish_for(start, dur):
    steps = max(int(-(-dur // 1)) - 1, 0)  # 0-day milestones and 1-day tasks finish on start
    return work_days_after(start, steps)


def month_miss(d, m):
    diff = (d.month - m) % 12
    return min(diff, 12 - diff)


def pick(start_m, finish_m, dur, rule, prev=None):
    """Choose a start date whose start/finish months best match the export's month hints."""
    if rule == "chain":
        lo, hi = prev[1] + timedelta(days=1), prev[1] + timedelta(days=100)
    else:
        lo, hi = PROJECT_START, PROJECT_FINISH_LIMIT
    best, d = [], lo
    while d <= hi:
        if d.weekday() != 6:
            f = finish_for(d, dur)
            if rule == "chain" or f <= PROJECT_FINISH_LIMIT:
                cost = month_miss(d, start_m) * 3 + (month_miss(f, finish_m) if finish_m else 0)
                best.append((cost, d, f))
        d += timedelta(days=1)
    low = min(c[0] for c in best)
    pool = [c for c in best if c[0] == low]
    if rule == "latest":
        c = max(pool, key=lambda x: x[1])
    else:
        c = min(pool, key=lambda x: x[1])
    return c[1], c[2], c[0]


def assign_blocks():
    block = "ZOY"
    for r in ROWS:
        if r.level == 0 or r.name.startswith("SITE MOBILIZATION"):
            block = "ZOY"
        elif r.level == 1:
            block = "A" if r.name.startswith("A-WING") else "B" if r.name.startswith("B-WING") else "EXT"
        r.block = block
    for i, r in enumerate(ROWS):
        nxt = ROWS[i + 1] if i + 1 < len(ROWS) else None
        r.is_summary = bool(nxt and nxt.level > r.level)


def schedule_leaves():
    in_chain = set(CHAIN_A) | set(CHAIN_B)
    # 1) rows that already started: earliest date that fits the hints
    for r in ROWS:
        if not r.is_summary and r.sr not in in_chain and r.as_m:
            r.start, r.finish, _ = pick(r.s_m, r.f_m, r.dur, "earliest")
    # 2) slab cycles, one after another, seeded from the sub-structure of each wing
    for chain, seed in ((CHAIN_A, 7), (CHAIN_B, 78)):
        prev = (BY_SR[seed].start, BY_SR[seed].finish)
        for sr in chain:
            r = BY_SR[sr]
            r.start, r.finish, _ = pick(r.s_m, r.f_m, r.dur, "chain", prev)
            prev = (r.start, r.finish)
    # 3) everything else: as late as the project window allows
    for r in ROWS:
        if not r.is_summary and r.start is None:
            r.start, r.finish, _ = pick(r.s_m, r.f_m, r.dur, "latest")


def roll_up():
    for i in range(len(ROWS) - 1, -1, -1):
        r = ROWS[i]
        if not r.is_summary:
            continue
        kids = []
        for k in ROWS[i + 1:]:
            if k.level <= r.level:
                break
            if not k.is_summary:
                kids.append(k)
        r.start = min(k.start for k in kids)
        r.finish = max(k.finish for k in kids)


def baselines_and_actuals():
    for r in ROWS:
        # baseline = hinted month closest before (or just after) the current start
        if r.bs_m and r.bf_m:
            best = None
            for y in range(2025, 2030):
                for d in (1, 8, 15, 22):
                    bs = first_workday(date(y, r.bs_m, d))
                    if bs < PROJECT_START or not (r.start - timedelta(days=335) <= bs <= r.start + timedelta(days=31)):
                        continue
                    bf = finish_for(bs, max(r.dur, 0))
                    diff = (bf.month - r.bf_m) % 12
                    miss = min(diff, 12 - diff)
                    key = (miss, abs((r.start - bs).days))
                    if best is None or key < best[0]:
                        best = (key, bs, bf)
            if best:
                r.b_start, r.b_finish = best[1], best[2]
        if r.b_start is None:
            r.b_start, r.b_finish = r.start, r.finish
        if r.as_m:
            r.a_start = r.start if r.start <= TODAY else None
        if r.af_m and r.a_start:
            r.a_finish = min(r.finish, TODAY)
    # summary rows: actual start = earliest child actual start
    for i in range(len(ROWS) - 1, -1, -1):
        r = ROWS[i]
        if r.is_summary:
            kids = []
            for k in ROWS[i + 1:]:
                if k.level <= r.level:
                    break
                if not k.is_summary:
                    kids.append(k)
            r.b_start = min(k.b_start for k in kids)
            r.b_finish = max(k.b_finish for k in kids)
            starts = [k.a_start for k in kids if k.a_start]
            r.a_start = min(starts) if starts else None


def link(succ, pred, relation="auto"):
    s, p = BY_SR[succ], BY_SR[pred]
    if s.is_summary or p.is_summary or s.start is None or p.start is None:
        return  # dependencies connect work rows only; summary rows take their dates from their children
    if relation in ("auto", "FS") and s.start > p.finish:
        lag = max((s.start - p.finish).days - 1, 0)
        s.preds.append((pred, "FS", min(lag, 0) if lag < 3 else lag))
    elif s.start >= p.start:
        s.preds.append((pred, "SS", (s.start - p.start).days))


def build_network():
    for chain, block in ((CHAIN_A, "A"), (CHAIN_B, "B")):
        base = 4 if block == "A" else 75
        exc = base + 1
        sub = exc + 2 if block == "A" else exc + 2  # sub-structure row follows RCC work header
        link(exc, base, "FS")
        link(sub, exc, "FS")
        link(chain[0], sub, "FS")
        for a, b in zip(chain, chain[1:]):
            link(b, a)
        # aluform advance -> shell plan -> material delivery
        adv, shell = (15, 16) if block == "A" else (85, 86)
        deliv = 17 if block == "A" else 87
        link(shell, adv, "FS")
        link(deliv, shell, "FS")
    # Finishing and MEP trades follow the structure by floor progress and by each other.
    for off in (0, 71 - 0 + 0):
        pass
    for wing, o in (("A", 0), ("B", 70)):
        sr = lambda a: a + o if wing == "A" else a + o
        blockwork, plaster, ext_grind = 49 + o, 50 + o, 51 + o
        wp, plumb, tiling = 52 + o, 53 + o, 54 + o
        carp, fab, alu, elec, paint = 55 + o, 56 + o, 57 + o, 58 + o, 59 + o
        ceil, san, cp, vdp, fire, lift = 60 + o, 61 + o, 62 + o, 63 + o, 64 + o, 65 + o
        solar, acid, grout, snag, clean, elev, other, noc = 66 + o, 67 + o, 68 + o, 69 + o, 70 + o, 71 + o, 72 + o, 73 + o
        slab_anchor = 22 + o  # 10 SLAB (5TH FLOOR)
        link(blockwork, slab_anchor)
        for succ, pred in [(plaster, blockwork), (ext_grind, blockwork), (wp, blockwork), (plumb, plaster),
                           (elec, plaster), (alu, blockwork), (fab, blockwork), (carp, plaster),
                           (tiling, plumb), (ceil, tiling), (paint, ceil), (san, tiling), (cp, san),
                           (vdp, paint), (fire, plaster), (lift, blockwork), (solar, paint),
                           (acid, tiling), (grout, acid), (elev, ext_grind), (snag, paint),
                           (clean, snag), (other, snag), (noc, snag)]:
            link(succ, pred)
    # external works
    link(146, 145, "FS")
    for succ, pred in [(147, 146), (148, 147), (150, 145), (151, 150), (152, 151), (153, 152), (154, 153)]:
        link(succ, pred)
    # Any work row still without a predecessor hangs off the closest earlier row of its own area.
    leaves = [r for r in ROWS if not r.is_summary]
    for r in leaves:
        if r.preds or r.sr == 2:
            continue
        same = [p for p in leaves if p is not r and p.block in (r.block, "ZOY")
                and (p.start, p.sr) < (r.start, r.sr)]
        if r.block in ("A", "B"):
            same = [p for p in same if p.block == r.block or p.sr == 2]
        anchor = max(same, key=lambda p: (p.start, p.sr))
        r.preds.append((anchor.sr, "SS", (r.start - anchor.start).days))


def wbs_codes():
    counters, codes = [0] * 6, {}
    for r in ROWS:
        counters[r.level] += 1
        for l in range(r.level + 1, 6):
            counters[l] = 0
        codes[r.sr] = ".".join(str(c) for c in counters[1:r.level + 1]) if r.level else "0"
    return codes


def fmt_pred(r):
    if not r.preds:
        return ""
    parts = []
    for sr, kind, lag in r.preds:
        t = f"{sr}{kind}" if not (kind == "FS" and lag == 0) else f"{sr}"
        if lag:
            t += f"+{lag}d"
        parts.append(t)
    return ",".join(parts)


def validate():
    issues = []
    for r in ROWS:
        for sr, kind, lag in r.preds:
            p = BY_SR[sr]
            if kind == "FS" and r.start <= p.finish:
                issues.append(f"{r.sr}: FS from {sr} violated")
            if kind == "SS" and r.start < p.start:
                issues.append(f"{r.sr}: SS from {sr} violated")
        if r.finish < r.start:
            issues.append(f"{r.sr}: finish before start")
        if r.a_start and r.a_start > TODAY:
            issues.append(f"{r.sr}: actual start in future")
        if r.pct > 0 and not r.a_start:
            issues.append(f"{r.sr}: progress without actual start")
    order = {r.sr: i for i, r in enumerate(sorted(ROWS, key=lambda x: (x.start, x.sr)))}
    for r in ROWS:
        for sr, _, _ in r.preds:
            if order[sr] >= order[r.sr] and BY_SR[sr].start >= r.start:
                issues.append(f"{r.sr}: predecessor {sr} does not come earlier (possible cycle)")
    leaves = [r for r in ROWS if not r.is_summary]
    mm_start = sum(1 for r in leaves if r.start.month != r.s_m)
    mm_finish = sum(1 for r in leaves if r.finish.month != r.f_m)
    return issues, len(leaves), mm_start, mm_finish


def write_workbook(path):
    codes = wbs_codes()
    wb = Workbook()
    ws = wb.active
    ws.title = "Schedule"
    head = ["ID", "WBS", "Level", "Task Name", "Wing / Area", "Type", "Duration (days)", "Baseline Start",
            "Baseline Finish", "Start", "Finish", "Actual Start", "Actual Finish", "% Complete",
            "Predecessors", "Site Engineer"]
    ws.append(head)
    for r in ROWS:
        eng = SITE_ENGINEER.get(r.block, "") if r.block in SITE_ENGINEER else ""
        area = {"A": "A-Wing", "B": "B-Wing", "EXT": "External Development", "ZOY": "Project"}[r.block]
        ws.append([r.sr, codes[r.sr], r.level, ("    " * r.level) + r.name, area,
                   "Summary" if r.is_summary else ("Milestone" if r.dur == 0 else "Task"),
                   r.dur, r.b_start, r.b_finish, r.start, r.finish, r.a_start, r.a_finish, r.pct / 100,
                   fmt_pred(r), eng])
    fill = PatternFill("solid", fgColor="1F3A5F")
    for c in ws[1]:
        c.font, c.fill = Font(bold=True, color="FFFFFF"), fill
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in ws.iter_rows(min_row=2):
        r = BY_SR[row[0].value]
        for c in row[7:13]:
            c.number_format = "dd-mmm-yyyy"
        row[13].number_format = "0%"
        if r.is_summary:
            for c in row:
                c.font = Font(bold=True)
                c.fill = PatternFill("solid", fgColor="E8EEF5")
        ws.row_dimensions[row[0].row].outlineLevel = max(r.level - 1, 0)
    widths = [6, 8, 7, 62, 20, 11, 14, 15, 15, 15, 15, 15, 15, 12, 22, 28]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "E2"
    ws.auto_filter.ref = ws.dimensions

    issues, n_leaves, mm_s, mm_f = validate()
    notes = wb.create_sheet("Read Me")
    lines = [
        ("41 ZOY - Sample Schedule (MVP template)", True),
        ("", False),
        ("Purpose: template for the schedule upload. Keep these column names and the ID / Predecessors format.", False),
        ("", False),
        ("How the dates were built (assumptions):", True),
        ("- The MSP export had masked dates (month only), so real dates were rebuilt: durations on a 6-day week (Sunday off).", False),
        ("- Slab cycles run one after another per wing; trades start as late as the project window allows, then follow their predecessors.", False),
        (f"- 'Today' for actuals and % complete is {TODAY:%d-%b-%Y}. No actual date is after today.", False),
        ("- Baseline dates are the hinted baseline months nearest the current start; they are an estimate, not the signed-off baseline.", False),
        ("- Summary rows (Type = Summary) take dates from their children. % Complete on summary rows is copied from the export.", False),
        ("- Predecessors use MSP style: '12' = finish-to-start, '14SS+30d' = start-to-start with 30 days lag. IDs are the ID column.", False),
        ("- Predecessor links are a proposed network (the export had none). The planner must confirm them before baselining.", False),
        ("- Site Engineer comes from the 41 ZOY organisation chart. External works engineer is not yet appointed.", False),
        ("", False),
        ("Self-check on this file:", True),
        (f"- Rows: {len(ROWS)} ({n_leaves} work rows). Dependency violations: {len(issues)}.", False),
        (f"- Work rows whose rebuilt start month differs from the export's month: {mm_s}; finish month: {mm_f}.", False),
    ]
    for text, bold in lines:
        notes.append([text])
        notes.cell(notes.max_row, 1).font = Font(bold=bold, size=13 if bold and notes.max_row == 1 else 11)
    notes.column_dimensions["A"].width = 130
    wb.move_sheet("Read Me", offset=-1)
    wb.active = 1
    wb.save(path)
    return issues, n_leaves, mm_s, mm_f


def main():
    out = Path(sys.argv[1] if len(sys.argv) > 1 else "data/zoy/41_ZOY_schedule_sample.xlsx")
    out.parent.mkdir(parents=True, exist_ok=True)
    assign_blocks()
    schedule_leaves()
    roll_up()
    baselines_and_actuals()
    build_network()
    issues, n_leaves, mm_s, mm_f = write_workbook(out)
    print(f"wrote {out}")
    print(f"rows={len(ROWS)} work_rows={n_leaves} violations={len(issues)} start_month_diff={mm_s} finish_month_diff={mm_f}")
    for i in issues[:20]:
        print(" ", i)


if __name__ == "__main__":
    main()
